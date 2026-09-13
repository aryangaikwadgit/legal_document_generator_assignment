import re
import unicodedata

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from src.schemas import GeneratedDocument


class DocumentBuilder:

    def __init__(self):
        self.document = Document()

        section = self.document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        self.usable_width = (
            section.page_width - section.left_margin - section.right_margin
        )

        style = self.document.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)

    def sanitize(self, text):
        replacements = {
            "\u25cf": "",
            "\u2011": "-",
            "\u2013": "-",
            "\u2014": "-",
            "\u2018": "'",
            "\u2019": "'",
            "\u201c": '"',
            "\u201d": '"',
        }

        for bad, good in replacements.items():
            text = text.replace(bad, good)

        return unicodedata.normalize("NFKC", text).strip()

    def format_paragraph(
        self,
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        bold=False,
        space_after=4,
        line_spacing=1.0,
    ):
        paragraph.alignment = alignment
        paragraph.paragraph_format.space_after = Pt(space_after)
        paragraph.paragraph_format.line_spacing = line_spacing

        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.bold = bold

    def add_text(
        self,
        text,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        bold=False,
        space_after=4,
    ):
        paragraph = self.document.add_paragraph()
        paragraph.add_run(self.sanitize(text))
        self.format_paragraph(
            paragraph,
            alignment,
            bold,
            space_after,
        )
        return paragraph

    def force_fixed_layout(self, table):
        table.autofit = False

        tbl_pr = table._tbl.tblPr

        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tbl_pr.append(layout)

        tbl_w = OxmlElement("w:tblW")
        tbl_w.set(qn("w:type"), "dxa")
        tbl_w.set(qn("w:w"), str(int(self.usable_width / 635)))
        tbl_pr.append(tbl_w)

    def add_heading(self, section):
        for line in section.content.splitlines():
            if line.strip():
                self.add_text(
                    line,
                    WD_ALIGN_PARAGRAPH.CENTER,
                    True,
                    2,
                )

    def add_party_line(self, text):
        text = self.sanitize(text)

        match = re.match(
            r"^(.*?)(\s+\.\.\.\s*.+)$",
            text,
        )

        if match:
            left_text = match.group(1).strip()
            right_text = match.group(2).strip()
        else:
            left_text = text
            right_text = ""

        left_width = self.usable_width * 0.74
        right_width = self.usable_width - left_width

        table = self.document.add_table(
            rows=1,
            cols=2,
        )

        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        self.force_fixed_layout(table)

        left_cell = table.cell(0, 0)
        right_cell = table.cell(0, 1)

        left_cell.width = left_width
        right_cell.width = right_width

        left_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        right_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

        left_paragraph = left_cell.paragraphs[0]
        left_paragraph.add_run(left_text)
        self.format_paragraph(
            left_paragraph,
            WD_ALIGN_PARAGRAPH.LEFT,
            False,
            1,
        )

        if right_text:
            right_paragraph = right_cell.paragraphs[0]
            right_paragraph.add_run(right_text)
            self.format_paragraph(
                right_paragraph,
                WD_ALIGN_PARAGRAPH.RIGHT,
                False,
                1,
            )

    def add_cause_title(self, section):
        for line in section.content.splitlines():
            line = line.strip()

            if not line:
                continue

            if line.upper() == "VERSUS":
                self.add_text(
                    line,
                    WD_ALIGN_PARAGRAPH.CENTER,
                    True,
                    3,
                )
            else:
                self.add_party_line(line)

    def add_numbered_paragraph(self, section):
        paragraph = self.document.add_paragraph()

        number_run = paragraph.add_run(f"{section.number}. ")
        number_run.bold = True

        paragraph.add_run(self.sanitize(section.content))

        self.format_paragraph(
            paragraph,
            WD_ALIGN_PARAGRAPH.JUSTIFY,
            False,
            5,
            1.0,
        )

    def add_prayer_item(self, section):
        label = section.label or "a"
        content = self.sanitize(section.content)

        if content.startswith(f"({label})"):
            content = content[len(label) + 2 :].strip()

        paragraph = self.document.add_paragraph()

        label_run = paragraph.add_run(f"({label}) ")
        label_run.bold = True

        paragraph.add_run(content)

        self.format_paragraph(
            paragraph,
            WD_ALIGN_PARAGRAPH.LEFT,
            False,
            4,
            1.0,
        )

    def add_jurat(self, section):
        lines = [
            self.sanitize(line)
            for line in section.content.splitlines()
            if line.strip() and line.strip().upper() not in ["DEPONENT", "BEFORE ME"]
        ]

        for line in lines:
            self.add_text(
                line,
                WD_ALIGN_PARAGRAPH.CENTER,
                False,
                2,
            )

        self.add_text(
            "Before Me",
            WD_ALIGN_PARAGRAPH.LEFT,
            False,
            2,
        )

        paragraph = self.document.add_paragraph()
        paragraph.add_run("DEPONENT").bold = True

        self.format_paragraph(
            paragraph,
            WD_ALIGN_PARAGRAPH.RIGHT,
            True,
            4,
        )

    def add_verification(self, section):
        lines = [
            self.sanitize(line)
            for line in section.content.splitlines()
            if line.strip() and line.strip().upper() not in ["VERIFICATION", "DEPONENT"]
        ]

        self.add_text(
            "VERIFICATION",
            WD_ALIGN_PARAGRAPH.CENTER,
            True,
            5,
        )

        for line in lines:
            if line:
                self.add_text(
                    line,
                    WD_ALIGN_PARAGRAPH.LEFT,
                    False,
                    4,
                )

        paragraph = self.document.add_paragraph()
        paragraph.add_run("DEPONENT").bold = True

        self.format_paragraph(
            paragraph,
            WD_ALIGN_PARAGRAPH.RIGHT,
            True,
            4,
        )

    def add_section(self, section):
        if section.type == "heading":
            self.add_heading(section)

        elif section.type == "cause_title":
            self.add_cause_title(section)

        elif section.type == "affidavit_title":
            self.add_text(
                section.content,
                WD_ALIGN_PARAGRAPH.CENTER,
                True,
                8,
            )

        elif section.type == "deponent":
            self.add_text(
                section.content,
                WD_ALIGN_PARAGRAPH.LEFT,
                False,
                6,
            )

        elif section.type == "numbered_paragraph":
            self.add_numbered_paragraph(section)

        elif section.type == "prayer_heading":
            self.add_text(
                section.content,
                WD_ALIGN_PARAGRAPH.CENTER,
                True,
                6,
            )

        elif section.type == "prayer_intro":
            self.add_text(
                section.content,
                WD_ALIGN_PARAGRAPH.LEFT,
                False,
                5,
            )

        elif section.type == "prayer_item":
            self.add_prayer_item(section)

        elif section.type == "jurat":
            self.add_jurat(section)

        elif section.type == "verification":
            self.add_verification(section)

        elif section.type == "advocate":
            for line in section.content.splitlines():
                if line.strip():
                    self.add_text(
                        line,
                        WD_ALIGN_PARAGRAPH.LEFT,
                        False,
                        2,
                    )

    def build(self, generated_document: GeneratedDocument):
        self.document = Document()

        section = self.document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        self.usable_width = (
            section.page_width - section.left_margin - section.right_margin
        )

        style = self.document.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)

        for section_data in generated_document.sections:
            self.add_section(section_data)

        return self.document

    def save(self, output_path):
        self.document.save(output_path)
